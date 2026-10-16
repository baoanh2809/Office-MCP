"""Blocks → .docx. Pure functions: no network, no disk.

``DocxRenderer`` is shared with edit_docx append_blocks, so it adapts to a foreign document: it looks styles
up by name and falls back to direct formatting when a style is missing, and it allocates numbering ids after
the ones the document already uses.
"""

import io

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import OxmlElement
from docx.parts.numbering import NumberingPart
from docx.shared import Cm, Emu, Inches, Mm, Pt

from app.core.safety import ImageInfo
from app.schemas.common import cell_text
from app.schemas.docx_spec import (
    BulletListBlock,
    HeadingBlock,
    ImageBlock,
    NumberedListBlock,
    PageBreakBlock,
    ParagraphBlock,
    TableBlock,
)

Images = dict[str, tuple[bytes, ImageInfo]]

EMU_PER_PIXEL = 9525  # 96 dpi
_HEADING_FALLBACK_PT = {1: 16, 2: 14, 3: 12, 4: 11}
_BORDERS = (
    "<w:tblBorders {}>"
    + "".join(
        f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="A6A6A6"/>'
        for side in ("top", "left", "bottom", "right", "insideH", "insideV")
    )
    + "</w:tblBorders>"
)


class DocxRenderer:
    def __init__(self, doc, images: Images) -> None:
        self.doc = doc
        self.images = images
        self._abstract_ids: dict[str, int] = {}
        self._bullet_num_id: int | None = None
        self._numbering = None

    # --- styles -------------------------------------------------------------
    def _style(self, name: str):
        try:
            return self.doc.styles[name]
        except KeyError:
            return None

    # --- numbering ----------------------------------------------------------
    def _numbering_el(self):
        if self._numbering is None:
            doc_part = self.doc.part
            try:
                part = doc_part.part_related_by(RT.NUMBERING)
            except KeyError:
                # python-docx cannot create a numbering part itself (NumberingPart.new is not implemented).
                package = doc_part.package
                part = NumberingPart(
                    package.next_partname("/word/numbering%d.xml"),
                    CT.WML_NUMBERING,
                    parse_xml(f"<w:numbering {nsdecls('w')}/>"),
                    package,
                )
                doc_part.relate_to(part, RT.NUMBERING)
            self._numbering = part.element
        return self._numbering

    def _ids(self, tag: str, attr: str) -> list[int]:
        out = []
        for el in self._numbering_el().findall(qn(tag)):
            try:
                out.append(int(el.get(qn(attr))))
            except (TypeError, ValueError):
                continue
        return out

    def _abstract(self, kind: str) -> int:
        if kind in self._abstract_ids:
            return self._abstract_ids[kind]
        new_id = max(self._ids("w:abstractNum", "w:abstractNumId"), default=-1) + 1
        fmt, text = ("bullet", "•") if kind == "bullet" else ("decimal", "%1.")
        el = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{new_id}"><w:multiLevelType w:val="singleLevel"/>'
            f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/>'
            '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl></w:abstractNum>'
        )
        numbering = self._numbering_el()
        nums = numbering.findall(qn("w:num"))
        abstracts = numbering.findall(qn("w:abstractNum"))
        if nums:
            nums[0].addprevious(el)  # every abstractNum must precede every num
        elif abstracts:
            abstracts[-1].addnext(el)
        else:
            numbering.insert(0, el)
        self._abstract_ids[kind] = new_id
        return new_id

    def _new_num(self, abstract_id: int, restart: bool) -> int:
        num_id = max(self._ids("w:num", "w:numId"), default=0) + 1
        override = '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>' if restart else ""
        el = parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>'
        )
        numbering = self._numbering_el()
        nums = numbering.findall(qn("w:num"))
        if nums:
            nums[-1].addnext(el)
        else:
            cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
            cleanup.addprevious(el) if cleanup is not None else numbering.append(el)
        return num_id

    # --- blocks -------------------------------------------------------------
    def render(self, blocks: list) -> None:
        for b in blocks:
            if isinstance(b, HeadingBlock):
                self._heading(b)
            elif isinstance(b, ParagraphBlock):
                run = self.doc.add_paragraph().add_run(b.text)
                run.bold = b.bold or None
                run.italic = b.italic or None
            elif isinstance(b, BulletListBlock):
                if self._bullet_num_id is None:
                    self._bullet_num_id = self._new_num(self._abstract("bullet"), restart=False)
                self._list(b.items, self._bullet_num_id)
            elif isinstance(b, NumberedListBlock):
                self._list(b.items, self._new_num(self._abstract("decimal"), restart=True))
            elif isinstance(b, TableBlock):
                self._table(b)
            elif isinstance(b, PageBreakBlock):
                self.doc.add_page_break()
            elif isinstance(b, ImageBlock):
                self._image(b)

    def _heading(self, b: HeadingBlock) -> None:
        style = self._style(f"Heading {b.level}")
        if style is not None:
            self.doc.add_paragraph(style=style).add_run(b.text)
            return
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(b.text)
        run.bold = True
        run.font.size = Pt(_HEADING_FALLBACK_PT[b.level])

    def _list(self, items: list[str], num_id: int) -> None:
        style = self._style("List Paragraph")
        for item in items:
            p = self.doc.add_paragraph(style=style) if style is not None else self.doc.add_paragraph()
            num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num_id
            p.add_run(item)

    def _table(self, b: TableBlock) -> None:
        cols = len(b.headers)
        table = self.doc.add_table(rows=len(b.rows) + 1, cols=cols)
        style = self._style("Table Grid")
        if style is not None:
            table.style = style
        else:
            table._tbl.tblPr.append(parse_xml(_BORDERS.format(nsdecls("w"))))
        header = table.rows[0]
        header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, text in zip(header.cells, b.headers, strict=True):
            cell.paragraphs[0].add_run(text).bold = True
        for r, row in enumerate(b.rows, start=1):
            cells = table.rows[r].cells
            for c in range(cols):
                text = cell_text(row[c]) if c < len(row) else ""
                if text:
                    cells[c].paragraphs[0].add_run(text)

    def text_width(self) -> int:
        # A foreign sectPr may omit page size or margins; fall back to Word's defaults (A4, 1-inch margins).
        section = self.doc.sections[-1]
        page = section.page_width or Mm(210)
        left = section.left_margin if section.left_margin is not None else Inches(1)
        right = section.right_margin if section.right_margin is not None else Inches(1)
        return int(page - left - right)

    def _image(self, b: ImageBlock) -> None:
        data, info = self.images[b.attachment_id]
        if b.width_cm is not None:
            width = Cm(b.width_cm)
        else:
            width = Emu(min(info.width * EMU_PER_PIXEL, self.text_width()))
        self.doc.add_paragraph().add_run().add_picture(io.BytesIO(data), width=width)


def new_document():
    """A blank document with a clean default look: A4, 2.54 cm margins, no template metadata."""
    doc = Document()
    package = doc.part.package
    for r_id, rel in list(package.rels.items()):
        if rel.reltype == RT.THUMBNAIL:
            package.rels.pop(r_id)
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.title = ""
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.54))
    return doc


def build_docx(blocks: list, images: Images) -> bytes:
    doc = new_document()
    DocxRenderer(doc, images).render(blocks)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
