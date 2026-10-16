"""Slides → 16:9 .pptx using python-pptx's default template.

The default template is 4:3, so its master and layout placeholders are widened to 16:9 before slides are
added. Placeholders a slide does not use are removed so no "Click to add text" boxes remain.
"""

import io

from pptx import Presentation
from pptx.util import Emu

from app.core.safety import ImageInfo
from app.schemas.common import cell_text
from app.schemas.pptx_spec import SlideSpec

Images = dict[str, tuple[bytes, ImageInfo]]

SLIDE_W, SLIDE_H = 12192000, 6858000  # 13.333 x 7.5 in
MARGIN = 457200
CONTENT_W = SLIDE_W - 2 * MARGIN
BODY_Y, BODY_H = 1371600, 5029200
ROW_H = 370840

# Default-template layout index per layout name.
_LAYOUT = {"title": 0, "title_and_content": 1, "section": 2, "two_content": 3, "table": 5, "image": 5}


def _has_own_xfrm(shape) -> bool:
    el = shape._element
    return bool(el.xpath("./p:spPr/a:xfrm") or el.xpath("./p:xfrm"))


def _widen_template(prs, factor: float) -> None:
    master = prs.slide_master
    targets = list(master.shapes) + [s for layout in master.slide_layouts for s in layout.shapes]
    for shape in targets:
        if _has_own_xfrm(shape):
            shape.left = Emu(int(shape.left * factor))
            shape.width = Emu(int(shape.width * factor))


def _set_lines(text_frame, items: list[str]) -> None:
    for i, item in enumerate(items):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = item.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\v")  # \v is a line break


def _fit(info: ImageInfo) -> tuple[int, int, int, int]:
    cx, cy = CONTENT_W, CONTENT_W * info.height // info.width
    if cy > BODY_H:
        cy, cx = BODY_H, BODY_H * info.width // info.height
    return MARGIN + (CONTENT_W - cx) // 2, BODY_Y + (BODY_H - cy) // 2, cx, cy


def _add_slide(prs, s: SlideSpec, images: Images) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[_LAYOUT[s.layout]])
    placeholders = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    used = {0}
    placeholders[0].text_frame.text = s.title
    if s.layout in ("title", "section") and s.subtitle is not None:
        _set_lines(placeholders[1].text_frame, [s.subtitle])
        used.add(1)
    if s.layout in ("title_and_content", "two_content") and s.bullets:
        _set_lines(placeholders[1].text_frame, s.bullets)
        used.add(1)
    if s.layout == "two_content" and s.right_bullets:
        _set_lines(placeholders[2].text_frame, s.right_bullets)
        used.add(2)
    for idx, ph in placeholders.items():
        if idx not in used:
            ph._element.getparent().remove(ph._element)
    if s.layout == "table":
        t = s.table
        shape = slide.shapes.add_table(
            len(t.rows) + 1, len(t.headers), Emu(MARGIN), Emu(BODY_Y), Emu(CONTENT_W), Emu(ROW_H * (len(t.rows) + 1))
        )
        table = shape.table
        for c, header in enumerate(t.headers):
            table.cell(0, c).text = header
        for r, row in enumerate(t.rows, start=1):
            for c in range(len(t.headers)):
                table.cell(r, c).text = cell_text(row[c]) if c < len(row) else ""
    if s.layout == "image":
        data, info = images[s.image_attachment_id]
        left, top, cx, cy = _fit(info)
        slide.shapes.add_picture(io.BytesIO(data), Emu(left), Emu(top), Emu(cx), Emu(cy))
    if s.notes is not None:
        slide.notes_slide.notes_text_frame.text = s.notes.replace("\r\n", "\n")


def build_pptx(slides: list[SlideSpec], images: Images) -> bytes:
    prs = Presentation()
    factor = SLIDE_W / prs.slide_width
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    _widen_template(prs, factor)
    props = prs.core_properties
    props.author = props.last_modified_by = props.title = ""
    for s in slides:
        _add_slide(prs, s, images)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
