"""Sheets → .xlsx with openpyxl's write-only mode (streams rows, up to 1 000 000 cells).

Bold, filled header row; frozen header; autofilter over the data; column widths from content; Excel
recalculates formulas on open (fullCalcOnLoad).
"""

import io

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.xlsx_spec import FormulaCell, SheetSpec, parse_date

_NUMBER_FORMAT = {"integer": "0", "currency": "#,##0.00", "percent": "0.00%", "date": "yyyy-mm-dd"}
_DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
_MIN_WIDTH, _MAX_WIDTH = 8, 60
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", start_color="D9E1F2")
_HEADER_BORDER = Border(bottom=Side(style="thin", color="8EA9DB"))


def text_cell(ws, value: str) -> Cell:
    """A cell that is always text: openpyxl would otherwise turn a leading "=" into a formula."""
    cell = WriteOnlyCell(ws, value=value)
    cell.data_type = "s"
    return cell


def _cell(ws, value, col_type: str) -> Cell | None:
    if value is None:
        return None
    if isinstance(value, FormulaCell):
        cell = WriteOnlyCell(ws, value=value.formula)
        if col_type in _NUMBER_FORMAT:
            cell.number_format = _NUMBER_FORMAT[col_type]
        return cell
    if col_type == "string":
        text = ("TRUE" if value else "FALSE") if isinstance(value, bool) else str(value)
        return text_cell(ws, text)
    if col_type == "date":
        parsed, has_time = parse_date(value)
        cell = WriteOnlyCell(ws, value=parsed)
        cell.number_format = _DATETIME_FORMAT if has_time else _NUMBER_FORMAT["date"]
        return cell
    if col_type == "integer":
        value = int(value)
    cell = WriteOnlyCell(ws, value=value)
    if col_type in _NUMBER_FORMAT:
        cell.number_format = _NUMBER_FORMAT[col_type]
    return cell


def _widths(sheet: SheetSpec) -> list[float]:
    widths = []
    for j, col in enumerate(sheet.columns):
        n = len(col.header)
        for row in sheet.rows:
            if j < len(row) and row[j] is not None and not isinstance(row[j], FormulaCell):
                n = max(n, len(str(row[j])))
        if col.type == "date":
            n = max(n, 19)
        widths.append(min(max(n, _MIN_WIDTH), _MAX_WIDTH) + 2)
    return widths


def build_xlsx(sheets: list[SheetSpec]) -> bytes:
    wb = Workbook(write_only=True)
    for sheet in sheets:
        ws = wb.create_sheet(title=sheet.name)
        for j, width in enumerate(_widths(sheet), start=1):
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{len(sheet.rows) + 1}"
        header = []
        for col in sheet.columns:
            cell = text_cell(ws, col.header)
            cell.font, cell.fill, cell.border = _HEADER_FONT, _HEADER_FILL, _HEADER_BORDER
            header.append(cell)
        ws.append(header)
        for row in sheet.rows:
            ws.append([_cell(ws, value, sheet.columns[c].type) for c, value in enumerate(row)])
    wb.calculation.fullCalcOnLoad = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
