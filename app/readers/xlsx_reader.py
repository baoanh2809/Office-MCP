"""read_xlsx: sheet list plus the values of one sheet range.

The workbook is opened twice in read-only (streaming) mode: once for formulas, once for Excel's cached
values. A formula cell returns ``{"formula": "=...", "value": <cached>}``; the cached value is null when the
file was never recalculated by Excel (for example a file written by a library). Date-formatted numbers
return ISO strings.
"""

import datetime as dt
import io
import re
import warnings

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from app.core import limits
from app.core.budget import TextBudget
from app.core.exceptions import OfficeDocsError, corrupt_file, validation_error

_RANGE = re.compile(r"^([A-Za-z]{1,3}[1-9][0-9]{0,6})(?::([A-Za-z]{1,3}[1-9][0-9]{0,6}))?$")


def _bounds(range_ref: str) -> tuple[int, int, int, int]:
    """(min_col, min_row, max_col, max_row) for "A1:D20" or "B2"."""
    bad = validation_error('range must be an A1 range such as "A1:D20" or one cell such as "B2"')
    m = _RANGE.match(range_ref.strip())
    if not m:
        raise bad
    try:
        c1, r1, c2, r2 = range_boundaries(f"{m.group(1)}:{m.group(2) or m.group(1)}".upper())
    except (ValueError, TypeError):
        raise bad from None
    if c2 < c1 or r2 < r1 or c2 > limits.MAX_EXCEL_COLUMNS or r2 > limits.MAX_EXCEL_ROWS:
        raise bad
    return c1, r1, c2, r2


def _ref(c1: int, r1: int, c2: int, r2: int) -> str:
    return f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"


def _scalar(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(timespec="seconds")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return str(value)
    if isinstance(value, float) and value.is_integer() and abs(value) < 2**53:
        return int(value)
    return value


def _open(data: bytes, data_only: bool):
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return load_workbook(io.BytesIO(data), read_only=True, data_only=data_only)
    except Exception:
        raise corrupt_file() from None


def _extent(ws) -> tuple[int, int] | None:
    """(max_col, max_row) from the dimension tag, else by scanning."""
    try:
        dim = ws.calculate_dimension()
        _, _, max_col, max_row = range_boundaries(dim)
        return max_col, max_row
    except (ValueError, TypeError):
        pass
    max_col = max_row = 0
    for row in ws.iter_rows():
        for cell in row:
            if getattr(cell, "value", None) is not None:
                max_col, max_row = max(max_col, cell.column), max(max_row, cell.row)
    return (max_col, max_row) if max_row else None


def read_xlsx(
    data: bytes,
    sheet: str | None = None,
    range_ref: str | None = None,
    max_rows: int | None = None,
    max_chars: int | None = None,
) -> dict:
    formulas_wb = _open(data, data_only=False)
    values_wb = _open(data, data_only=True)
    try:
        return _read(formulas_wb, values_wb, sheet, range_ref, max_rows, max_chars)
    except OfficeDocsError:
        raise
    except Exception:
        raise corrupt_file() from None
    finally:
        formulas_wb.close()
        values_wb.close()


def _read(formulas_wb, values_wb, sheet, range_ref, max_rows, max_chars) -> dict:
    names = formulas_wb.sheetnames
    if not names:
        raise corrupt_file()
    infos = []
    for ws in formulas_wb.worksheets:
        try:
            dimension = ws.calculate_dimension()
        except (ValueError, TypeError):
            dimension = None
        infos.append({"name": ws.title, "dimension": dimension, "visible": ws.sheet_state == "visible"})
    name = names[0]
    if sheet:
        if sheet not in names:
            raise validation_error("sheet does not exist in this workbook; call read_xlsx without sheet to list sheets")
        name = sheet
    ws_f, ws_v = formulas_wb[name], values_wb[name]
    result = {"sheets": infos, "sheet": name, "range": None, "rows": [], "row_count": 0, "truncated": False}

    requested = _bounds(range_ref) if range_ref else None
    extent = _extent(ws_f)
    if requested:
        c1, r1, c2, r2 = requested
        result["range"] = _ref(*requested)
    elif extent:
        c1, r1, c2, r2 = 1, 1, *extent
        result["range"] = _ref(c1, r1, c2, r2)
    else:
        return result
    if extent:
        r2 = min(r2, extent[1])  # do not walk empty rows past the data

    row_cap = limits.clamp(max_rows, limits.DEFAULT_MAX_ROWS, limits.MAX_MAX_ROWS)
    budget = TextBudget(limits.clamp(max_chars, limits.DEFAULT_MAX_CHARS, limits.MAX_MAX_CHARS))
    if r2 < r1:
        return result
    rows_f = ws_f.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2)
    rows_v = ws_v.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2)
    for row_f, row_v in zip(rows_f, rows_v, strict=False):
        if len(result["rows"]) == row_cap:
            result["truncated"] = True
            break
        out, complete = [], True
        for cell_f, cell_v in zip(row_f, row_v, strict=False):
            value, cost = _convert(cell_f, cell_v)
            if not budget.spend(cost):
                complete = False
                break
            out.append(value)
        width = c2 - c1 + 1
        if complete:
            out.extend([None] * (width - len(out)))
        result["rows"].append(out)
        if not complete:
            result["truncated"] = True
            break
    result["row_count"] = len(result["rows"])
    return result


def _convert(cell_f, cell_v):
    raw = getattr(cell_f, "value", None)
    if getattr(cell_f, "data_type", None) == "f" or isinstance(raw, ArrayFormula):
        formula = raw.text if isinstance(raw, ArrayFormula) else str(raw)
        if not formula.startswith("="):
            formula = "=" + formula
        cached = _scalar(getattr(cell_v, "value", None))
        return {"formula": formula, "value": cached}, len(formula) + len(str(cached if cached is not None else ""))
    if raw is None:
        return None, 0
    value = _scalar(raw)
    return value, len(str(value))
