"""read_pptx: slides in presentation order with title, text paragraphs, tables, and speaker notes."""

import io

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.group import GroupShape

from app.core import limits
from app.core.budget import TextBudget
from app.core.exceptions import OfficeDocsError, corrupt_file

_TITLES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
_CHROME = {
    PP_PLACEHOLDER.SLIDE_NUMBER,
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.HEADER,
    PP_PLACEHOLDER.SLIDE_IMAGE,
}


def _paragraphs(text_frame) -> list[str]:
    out = []
    for p in text_frame.paragraphs:
        text = p.text.replace("\v", "\n")  # python-pptx returns a:br as vertical tab
        if text.strip():
            out.append(text)
    return out


def _placeholder_type(shape):
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except (ValueError, AttributeError):
        return None


def _collect(shapes, content: dict) -> None:
    for shape in shapes:
        if isinstance(shape, GroupShape):  # shape_type raises on shapes without a preset geometry
            _collect(shape.shapes, content)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [["\n".join(_paragraphs(cell.text_frame)) for cell in row.cells] for row in shape.table.rows]
            content["tables"].append(rows)
            continue
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            continue
        ph = _placeholder_type(shape)
        if ph in _CHROME:
            continue
        paragraphs = _paragraphs(shape.text_frame)
        if ph in _TITLES and content["title"] is None:
            content["title"] = "\n".join(paragraphs)
        else:
            content["texts"].extend(paragraphs)


def _notes(slide) -> str | None:
    if not slide.has_notes_slide:
        return None
    ph = slide.notes_slide.notes_placeholder
    if ph is None:
        return None
    text = "\n".join(_paragraphs(ph.text_frame))
    return text or None


def read_pptx(data: bytes, max_chars: int | None = None) -> dict:
    try:
        prs = Presentation(io.BytesIO(data))
        slides = list(prs.slides)
    except Exception:
        raise corrupt_file() from None
    budget = TextBudget(limits.clamp(max_chars, limits.DEFAULT_MAX_CHARS, limits.MAX_MAX_CHARS))
    out: list[dict] = []
    try:
        for index, slide in enumerate(slides, start=1):
            if budget.truncated:
                break
            out.append(_read_slide(slide, index, budget))
    except OfficeDocsError:
        raise
    except Exception:
        raise corrupt_file() from None
    return {"slide_count": len(slides), "slides": out, "truncated": budget.truncated}


def _read_slide(slide, index: int, budget: TextBudget) -> dict:
    content = {"title": None, "texts": [], "tables": []}
    _collect(slide.shapes, content)
    result = {"index": index, "title": None, "texts": [], "tables": [], "notes": None}
    if content["title"] is not None:
        result["title"] = budget.take(content["title"])
    for text in content["texts"]:
        if budget.truncated:
            return result
        if taken := budget.take(text):
            result["texts"].append(taken)
    for table in content["tables"]:
        if budget.truncated:
            return result
        rows = []
        for row in table:
            cells = []
            for cell in row:
                if budget.exhausted and cell:
                    budget.truncated = True
                    break
                cells.append(budget.take(cell))
            rows.append(cells)
            if budget.truncated:
                break
        result["tables"].append({"rows": rows})
    if not budget.truncated:
        notes = _notes(slide)
        if notes is not None:
            result["notes"] = budget.take(notes)
    return result
