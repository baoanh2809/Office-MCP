"""create_docx builder, read_docx reader, edit_docx editor."""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree
from PIL import Image
from pydantic import TypeAdapter, ValidationError

from app.builders.docx_builder import build_docx
from app.core import limits
from app.core.exceptions import ErrorCode, OfficeDocsError
from app.core.safety import check_image, inspect_package
from app.editors.docx_editor import edit_docx
from app.readers.docx_reader import read_docx
from app.schemas.docx_spec import DocxBlock, DocxOperation, image_ids, validate_blocks, validate_operations

FIXTURE = Path(__file__).parent / "fixtures" / "sample.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BLOCKS = TypeAdapter(list[DocxBlock])
OPS = TypeAdapter(list[DocxOperation])


def png(w: int, h: int) -> tuple[bytes, object]:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (0, 0, 255)).save(buf, format="PNG")
    return buf.getvalue(), check_image(buf.getvalue())


def part(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return zf.read(name)


def parts(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return zf.namelist()


def _code(fn, *args) -> ErrorCode:
    with pytest.raises(OfficeDocsError) as exc:
        fn(*args)
    return exc.value.code


EVERY_BLOCK = [
    {"type": "heading", "level": 1, "text": "Q3 report"},
    {"type": "heading", "level": 4, "text": "Small heading"},
    {"type": "paragraph", "text": "Plain & <safe> text"},
    {"type": "paragraph", "text": "Bold italic", "bold": True, "italic": True},
    {"type": "paragraph", "text": "Line one\nLine\ttwo"},
    {"type": "bullet_list", "items": ["Alpha", "Beta"]},
    {"type": "numbered_list", "items": ["One", "Two"]},
    {"type": "numbered_list", "items": ["Again one"]},
    {"type": "table", "headers": ["Region", "Total"], "rows": [["North", 1200], ["South"]]},
    {"type": "page_break"},
    {"type": "image", "attachment_id": "att_img", "width_cm": 5},
    {"type": "paragraph", "text": "After image"},
]


def build_every() -> bytes:
    blocks = BLOCKS.validate_python(EVERY_BLOCK)
    validate_blocks("blocks", blocks)
    return build_docx(blocks, {"att_img": png(40, 20)})


def test_every_block_reads_back():
    data = build_every()
    inspect_package(data, "docx")
    res = read_docx(data)
    assert res.blocks == [
        {"type": "heading", "level": 1, "text": "Q3 report"},
        {"type": "heading", "level": 4, "text": "Small heading"},
        {"type": "paragraph", "text": "Plain & <safe> text"},
        {"type": "paragraph", "text": "Bold italic"},
        {"type": "paragraph", "text": "Line one\nLine\ttwo"},
        {"type": "list", "ordered": False, "items": ["Alpha", "Beta"]},
        {"type": "list", "ordered": True, "items": ["One", "Two"]},
        {"type": "list", "ordered": True, "items": ["Again one"]},
        {"type": "table", "rows": [["Region", "Total"], ["North", "1200"], ["South", ""]]},
        {"type": "paragraph", "text": "After image"},
    ]
    assert not res.truncated


def test_reopened_with_python_docx():
    doc = Document(io.BytesIO(build_every()))
    styles = [p.style.name for p in doc.paragraphs[:2]]
    assert styles == ["Heading 1", "Heading 4"]
    bold_italic = doc.paragraphs[3].runs[0]
    assert bold_italic.bold and bold_italic.italic
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Region", "Total"]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold
    shape = doc.inline_shapes[0]
    assert shape.width == 1_800_000 and shape.height == 900_000  # 5 cm, 2:1 kept


def test_numbering_restarts_per_numbered_list():
    doc = etree.fromstring(part(build_every(), "word/document.xml"))
    ids = [n.get(W + "val") for n in doc.iter(W + "numId")]
    assert len(ids) == 5
    assert ids[0] == ids[1] and ids[2] == ids[3] and ids[3] != ids[4]
    numbering = part(build_every(), "word/numbering.xml")
    assert numbering.count(b"startOverride") == 2
    assert b"page" in part(build_every(), "word/document.xml")


def test_default_image_width_caps_at_text_width():
    data = build_docx(BLOCKS.validate_python([{"type": "image", "attachment_id": "big"}]), {"big": png(3000, 1500)})
    shape = Document(io.BytesIO(data)).inline_shapes[0]
    section = Document(io.BytesIO(data)).sections[0]
    assert shape.width == section.page_width - section.left_margin - section.right_margin
    assert shape.height == shape.width // 2


@pytest.mark.parametrize(
    "blocks",
    [
        [],
        [{"type": "heading", "level": 1, "text": "  "}],
        [{"type": "heading", "level": 1, "text": "a\nb"}],
        [{"type": "bullet_list", "items": []}],
        [{"type": "table", "headers": [], "rows": []}],
        [{"type": "table", "headers": ["a"], "rows": [["1", "2"]]}],
        [{"type": "image", "attachment_id": "../x"}],
        [{"type": "image", "attachment_id": "a", "width_cm": 0.1}],
        [{"type": "paragraph", "text": "a\x07b"}],
    ],
)
def test_validate_blocks_rejects(blocks):
    assert _code(validate_blocks, "blocks", BLOCKS.validate_python(blocks)) == ErrorCode.VALIDATION_ERROR


def test_too_many_blocks():
    blocks = BLOCKS.validate_python([{"type": "page_break"}] * (limits.MAX_BLOCKS + 1))
    assert _code(validate_blocks, "blocks", blocks) == ErrorCode.VALIDATION_ERROR


@pytest.mark.parametrize(
    "raw",
    [
        [{"type": "quote", "text": "x"}],
        [{"type": "heading", "text": "x"}],
        [{"type": "heading", "level": 5, "text": "x"}],
        [{"type": "paragraph", "text": "x", "items": ["a"]}],
        [{"type": "page_break", "text": "x"}],
    ],
)
def test_schema_rejects_wrong_shapes(raw):
    with pytest.raises(ValidationError):
        BLOCKS.validate_python(raw)


def test_image_ids_deduplicated():
    blocks = BLOCKS.validate_python(
        [
            {"type": "image", "attachment_id": "a"},
            {"type": "paragraph", "text": "x"},
            {"type": "image", "attachment_id": "a"},
            {"type": "image", "attachment_id": "b"},
        ]
    )
    assert image_ids(blocks) == ["a", "b"]


# --- reader ---------------------------------------------------------------------


def test_read_fixture():
    res = read_docx(FIXTURE.read_bytes())
    assert res.blocks == [
        {"type": "heading", "level": 1, "text": "Quarterly review"},
        {"type": "paragraph", "text": "Revenue for Acme Corp grew."},
        {"type": "paragraph", "text": "See the dashboard for details."},
        {"type": "list", "ordered": False, "items": ["First bullet", "Second bullet"]},
        {"type": "list", "ordered": True, "items": ["Step one", "Step two"]},
        {"type": "paragraph", "text": "Inside a content control & more."},
        {"type": "paragraph", "text": "Before box."},
        {"type": "table", "rows": [["Region", "Total"], ["North", "1,200\n(est.)"]]},
        {"type": "heading", "level": 2, "text": "Tab\tbed"},
    ]


def test_read_truncates():
    res = read_docx(FIXTURE.read_bytes(), max_chars=20)
    assert res.truncated and len(res.blocks) == 2 and res.char_count <= 20
    assert res.blocks[1]["text"] == "Reve"


def test_read_corrupt():
    assert _code(read_docx, b"PK\x03\x04nope") == ErrorCode.CORRUPT_FILE


# --- editor ---------------------------------------------------------------------


def edit(src: bytes, raw: list[dict], images=None):
    ops = OPS.validate_python(raw)
    validate_operations(ops)
    out, results = edit_docx(src, ops, images or {})
    inspect_package(out, "docx")
    return out, results


def test_replace_across_runs_and_headers():
    src = FIXTURE.read_bytes()
    before = bytes(src)
    out, results = edit(src, [{"type": "replace_text", "find": "Acme Corp", "replace": "Globex Inc"}])
    assert src == before
    assert results == [{"type": "replace_text", "replacements": 2}]
    assert read_docx(out).blocks[1]["text"] == "Revenue for Globex Inc grew."
    assert b"Globex Inc confidential" in part(out, "word/header1.xml")
    # The replacement takes the formatting of the run where the match starts (bold "Acme").
    first_bold = Document(io.BytesIO(out)).paragraphs[1].runs[1]
    assert first_bold.bold and first_bold.text == "Globex Inc"


@pytest.mark.parametrize(
    "op,count",
    [
        ({"type": "replace_text", "find": "QUARTERLY", "replace": "Annual", "match_case": False}, 1),
        ({"type": "replace_text", "find": "QUARTERLY", "replace": "Annual"}, 0),
        ({"type": "replace_text", "find": "TEXT BOX CONTENT", "replace": "boxed"}, 1),
        ({"type": "replace_text", "find": "Tabbed", "replace": "x"}, 0),
        ({"type": "replace_text", "find": "deleted text", "replace": "x"}, 0),
        ({"type": "replace_text", "find": "e", "replace": "E"}, 28),
    ],
)
def test_replace_options_and_boundaries(op, count):
    _, results = edit(FIXTURE.read_bytes(), [op])
    assert results[0]["replacements"] == count


def test_replace_preserves_spaces_and_escapes():
    out, _ = edit(FIXTURE.read_bytes(), [{"type": "replace_text", "find": "Before box.", "replace": " R&D <box> "}])
    assert read_docx(out).blocks[6]["text"] == " R&D <box> "
    assert b'xml:space="preserve"> R&amp;D &lt;box&gt; </w:t>' in part(out, "word/document.xml")


def test_append_blocks_to_foreign_document():
    out, results = edit(
        FIXTURE.read_bytes(),
        [
            {
                "type": "append_blocks",
                "blocks": [
                    {"type": "heading", "level": 1, "text": "Appendix"},
                    {"type": "numbered_list", "items": ["New one", "New two"]},
                    {"type": "table", "headers": ["k"], "rows": [["v"]]},
                    {"type": "image", "attachment_id": "att_img"},
                ],
            }
        ],
        {"att_img": png(10, 10)},
    )
    assert results == [{"type": "append_blocks", "blocks_added": 4}]
    blocks = read_docx(out).blocks
    assert blocks[-3:] == [
        {"type": "heading", "level": 1, "text": "Appendix"},
        {"type": "list", "ordered": True, "items": ["New one", "New two"]},
        {"type": "table", "rows": [["k"], ["v"]]},
    ]
    body = etree.fromstring(part(out, "word/document.xml")).find(W + "body")
    assert body[-1].tag == W + "sectPr", "sectPr must stay the last body element"
    heading = [p for p in body.iter(W + "p") if "".join(p.itertext()) == "Appendix"][0]
    assert heading.find(f"{W}pPr/{W}pStyle").get(W + "val") == "Titre1", "reuse the document's own heading style"
    numbering = etree.fromstring(part(out, "word/numbering.xml"))
    children = [c.tag.split("}")[1] for c in numbering]
    assert children.index("num") > max(i for i, t in enumerate(children) if t == "abstractNum")
    assert any(n.endswith(".png") for n in parts(out) if n.startswith("word/media/"))
    assert len(Document(io.BytesIO(out)).inline_shapes) == 1


def _without_numbering(data: bytes) -> bytes:
    """Drop word/numbering.xml, its relationship, and its content type."""
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name in src.namelist():
            if name == "word/numbering.xml":
                continue
            body = src.read(name)
            if name in ("word/_rels/document.xml.rels", "[Content_Types].xml"):
                root = etree.fromstring(body)
                for el in list(root):
                    if "numbering" in (el.get("Type", "") + el.get("PartName", "")):
                        root.remove(el)
                body = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            out.writestr(name, body)
    return buf.getvalue()


def test_append_list_creates_numbering_part_when_missing():
    src = _without_numbering(build_docx(BLOCKS.validate_python([{"type": "paragraph", "text": "only text"}]), {}))
    assert "word/numbering.xml" not in parts(src)
    out, _ = edit(src, [{"type": "append_blocks", "blocks": [{"type": "bullet_list", "items": ["a", "b"]}]}])
    assert any("numbering" in n for n in parts(out))
    assert read_docx(out).blocks[1] == {"type": "list", "ordered": False, "items": ["a", "b"]}


def test_builder_drops_template_thumbnail_and_author():
    data = build_every()
    assert not any("thumbnail" in n for n in parts(data))
    assert Document(io.BytesIO(data)).core_properties.author == ""


def test_edit_sequence_applies_in_order():
    out, results = edit(
        FIXTURE.read_bytes(),
        [
            {"type": "append_blocks", "blocks": [{"type": "paragraph", "text": "Draft by Acme Corp"}]},
            {"type": "replace_text", "find": "Acme Corp", "replace": "Initech"},
        ],
    )
    assert results[1]["replacements"] == 3
    assert read_docx(out).blocks[-1]["text"] == "Draft by Initech"


@pytest.mark.parametrize(
    "raw",
    [
        [],
        [{"type": "replace_text", "find": "", "replace": "x"}],
        [{"type": "replace_text", "find": "a\nb", "replace": "x"}],
        [{"type": "append_blocks", "blocks": []}],
    ],
)
def test_validate_operations_rejects(raw):
    assert _code(validate_operations, OPS.validate_python(raw)) == ErrorCode.VALIDATION_ERROR


@pytest.mark.parametrize(
    "raw",
    [
        [{"type": "delete_all"}],
        [{"type": "replace_text", "find": "x"}],
        [{"type": "replace_text", "find": "a", "replace": "b", "blocks": []}],
        [{"type": "append_blocks"}],
    ],
)
def test_operation_schema_rejects_wrong_shapes(raw):
    with pytest.raises(ValidationError):
        OPS.validate_python(raw)
