"""create_xlsx builder, read_xlsx reader, update_xlsx editor (formula injection included)."""

import datetime as dt
import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from pydantic import TypeAdapter, ValidationError

from app.builders.xlsx_builder import build_xlsx
from app.core.exceptions import ErrorCode, OfficeDocsError
from app.core.safety import inspect_package
from app.editors.xlsx_editor import update_xlsx
from app.readers.xlsx_reader import read_xlsx
from app.schemas.xlsx_spec import AppendRows, CellUpdate, SheetSpec, validate_sheets, validate_update

FIXTURE = Path(__file__).parent / "fixtures" / "sample.xlsx"
SHEETS = TypeAdapter(list[SheetSpec])
UPDATES = TypeAdapter(list[CellUpdate])
APPENDS = TypeAdapter(list[AppendRows])


def _code(fn, *args) -> ErrorCode:
    with pytest.raises(OfficeDocsError) as exc:
        fn(*args)
    return exc.value.code


ALL_TYPES = [
    {
        "name": "Sales",
        "columns": [
            {"header": "Item", "type": "string"},
            {"header": "Qty", "type": "integer"},
            {"header": "Price", "type": "currency"},
            {"header": "Ratio", "type": "number"},
            {"header": "Share", "type": "percent"},
            {"header": "Day", "type": "date"},
            {"header": "Paid", "type": "boolean"},
        ],
        "rows": [
            ["Pens", 3, 1.5, 0.333, 0.25, "2026-01-05", True],
            ["Paper", 4, 2, 1e3, 0.5, "2026-01-06T13:30:00", False],
            [None, {"formula": "=SUM(B2:B3)"}, {"formula": "=SUM(C2:C3)"}],
        ],
    },
    {
        "name": "Notes",
        "columns": [{"header": "Text", "type": "string"}],
        "rows": [['=HYPERLINK("http://x")'], ["+1"], ["-1"], ["@SUM(A1)"], ["\tcmd"], [42]],
    },
]


def build(raw=ALL_TYPES) -> bytes:
    sheets = SHEETS.validate_python(raw)
    validate_sheets(sheets)
    data = build_xlsx(sheets)
    inspect_package(data, "xlsx")
    return data


def test_every_column_type_is_written():
    wb = load_workbook(io.BytesIO(build()))
    assert wb.sheetnames == ["Sales", "Notes"]
    ws = wb["Sales"]
    assert [c.value for c in ws[1]] == ["Item", "Qty", "Price", "Ratio", "Share", "Day", "Paid"]
    assert [c.value for c in ws[2]] == ["Pens", 3, 1.5, 0.333, 0.25, dt.datetime(2026, 1, 5), True]
    assert ws["D3"].value == 1000 and ws["G3"].value is False
    assert ws["F3"].value == dt.datetime(2026, 1, 6, 13, 30)
    assert ws["F2"].number_format == "yyyy-mm-dd" and ws["F3"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert ws["C2"].number_format == "#,##0.00" and ws["E2"].number_format == "0.00%"
    assert ws["B2"].number_format == "0"
    assert ws["B4"].value == "=SUM(B2:B3)" and ws["B4"].data_type == "f"


def test_header_freeze_filter_widths_and_recalc():
    wb = load_workbook(io.BytesIO(build()))
    ws = wb["Sales"]
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:G4"
    assert ws.column_dimensions["A"].width > 8
    assert wb.calculation.fullCalcOnLoad


def test_formula_like_strings_are_written_as_text():
    ws = load_workbook(io.BytesIO(build()))["Notes"]
    expected = ['=HYPERLINK("http://x")', "+1", "-1", "@SUM(A1)", "\tcmd", "42"]
    for row, want in enumerate(expected, start=2):
        cell = ws.cell(row=row, column=1)
        assert (cell.value, cell.data_type) == (want, "s"), row


@pytest.mark.parametrize(
    "raw",
    [
        [],
        [
            {"name": "A", "columns": [{"header": "x", "type": "string"}]},
            {"name": "a", "columns": [{"header": "x", "type": "string"}]},
        ],
        [{"name": "Q1/Q2", "columns": [{"header": "x", "type": "string"}]}],
        [{"name": "n" * 32, "columns": [{"header": "x", "type": "string"}]}],
        [{"name": "History", "columns": [{"header": "x", "type": "string"}]}],
        [{"name": "A", "columns": []}],
        [{"name": "A", "columns": [{"header": "x", "type": "string"}], "rows": [["a", "b"]]}],
        [{"name": "A", "columns": [{"header": "x", "type": "number"}], "rows": [["12"]]}],
        [{"name": "A", "columns": [{"header": "x", "type": "integer"}], "rows": [[1.5]]}],
        [{"name": "A", "columns": [{"header": "x", "type": "date"}], "rows": [["05/01/2026"]]}],
        [{"name": "A", "columns": [{"header": "x", "type": "boolean"}], "rows": [[1]]}],
        [
            {
                "name": "A",
                "columns": [{"header": "x", "type": "number"}],
                "rows": [[{"formula": '=WEBSERVICE("http://x")'}]],
            }
        ],
        [{"name": "A", "columns": [{"header": "x", "type": "number"}], "rows": [[{"formula": "SUM(A1)"}]]}],
        [{"name": "A", "columns": [{"header": "x", "type": "string"}], "rows": [["a\x01b"]]}],
    ],
)
def test_validate_sheets_rejects(raw):
    assert _code(validate_sheets, SHEETS.validate_python(raw)) == ErrorCode.VALIDATION_ERROR


def test_cell_limit():
    columns = [{"header": "c", "type": "string"}] * 1000
    sheets = SHEETS.validate_python([{"name": "Big", "columns": columns, "rows": [[]] * 1000}])
    with pytest.raises(OfficeDocsError) as exc:
        validate_sheets(sheets)
    assert exc.value.code == ErrorCode.VALIDATION_ERROR and "1000000" in exc.value.message


@pytest.mark.parametrize("cell", [{"value": 1}, {"formula": "=1", "x": 2}, [1], {"formula": 5}])
def test_schema_rejects_cell_shapes(cell):
    with pytest.raises(ValidationError):
        SHEETS.validate_python([{"name": "A", "columns": [{"header": "x", "type": "string"}], "rows": [[cell]]}])


def test_schema_rejects_unknown_column_type():
    with pytest.raises(ValidationError):
        SHEETS.validate_python([{"name": "A", "columns": [{"header": "x", "type": "money"}]}])


# --- reader ---------------------------------------------------------------------


def test_read_fixture():
    res = read_xlsx(FIXTURE.read_bytes())
    assert res["sheets"] == [
        {"name": "Data", "dimension": "A1:D5", "visible": True},
        {"name": "Hidden", "dimension": None, "visible": False},
    ]
    assert (res["sheet"], res["range"], res["truncated"], res["row_count"]) == ("Data", "A1:D5", False, 5)
    assert res["rows"] == [
        ["Item", "Amount", "Date", None],
        ["Pens", 3, "2024-01-01", 0.25],
        ["Paper", 4.5, None, None],
        ["Total", {"formula": "=SUM(B2:B3)", "value": 7.5}, {"formula": "=B4*2", "value": None}, None],
        ["=not a formula", True, None, None],
    ]


def test_read_range_and_sheet():
    res = read_xlsx(FIXTURE.read_bytes(), range_ref="b2:c3")
    assert (res["range"], res["rows"]) == ("B2:C3", [[3, "2024-01-01"], [4.5, None]])
    res = read_xlsx(FIXTURE.read_bytes(), sheet="Hidden")
    assert (res["range"], res["rows"]) == ("A1:A1", [["secret sheet text"]])


@pytest.mark.parametrize(
    "kwargs",
    [{"sheet": "Nope"}, {"range_ref": "A0:B2"}, {"range_ref": "ZZZZ1"}, {"range_ref": "B2:A1"}, {"range_ref": "1:2"}],
)
def test_read_rejects_bad_sheet_or_range(kwargs):
    assert _code(lambda: read_xlsx(FIXTURE.read_bytes(), **kwargs)) == ErrorCode.VALIDATION_ERROR


def test_read_truncation():
    res = read_xlsx(FIXTURE.read_bytes(), max_rows=2)
    assert (len(res["rows"]), res["row_count"], res["truncated"]) == (2, 2, True)
    res = read_xlsx(FIXTURE.read_bytes(), max_chars=12)
    # The row stops at the first cell that does not fit; no null padding that would look like empty cells.
    assert res["truncated"] and res["rows"] == [["Item", "Amount"]]


def test_read_builder_output():
    res = read_xlsx(build())
    assert res["rows"][1] == ["Pens", 3, 1.5, 0.333, 0.25, "2026-01-05", True]
    assert res["rows"][2][5] == "2026-01-06T13:30:00"
    assert res["rows"][3][1] == {"formula": "=SUM(B2:B3)", "value": None}


def test_read_corrupt():
    assert _code(read_xlsx, b"PK\x03\x04garbage") == ErrorCode.CORRUPT_FILE


# --- update ---------------------------------------------------------------------


def test_update_cells_and_append_rows():
    src = FIXTURE.read_bytes()
    before = bytes(src)
    updates = UPDATES.validate_python(
        [
            {"sheet": "Data", "cell": "B2", "value": 10},
            {"sheet": "Data", "cell": "C3", "formula": "=B3*2"},
            {"sheet": "Data", "cell": "A3", "value": "=1+1"},
            {"sheet": "Data", "cell": "D2", "value": None},
        ]
    )
    appends = APPENDS.validate_python([{"sheet": "Data", "rows": [["Ink", 2.25], ["Sum", {"formula": "=SUM(B2:B6)"}]]}])
    validate_update(updates, appends)
    out, res = update_xlsx(src, updates, appends)
    assert src == before
    assert (res["cells_updated"], res["rows_appended"]) == (4, 2)
    inspect_package(out, "xlsx")
    wb = load_workbook(io.BytesIO(out))
    ws = wb["Data"]
    assert ws["B2"].value == 10 and ws["D2"].value is None
    assert (ws["A3"].value, ws["A3"].data_type) == ("=1+1", "s")
    assert ws["C3"].value == "=B3*2"
    assert [ws["A6"].value, ws["B6"].value, ws["A7"].value, ws["B7"].value] == ["Ink", 2.25, "Sum", "=SUM(B2:B6)"]
    assert wb.calculation.fullCalcOnLoad


@pytest.mark.parametrize(
    "updates",
    [
        [{"sheet": "Data", "cell": "A1", "value": 1, "formula": "=1"}],
        [{"sheet": "Data", "cell": "A1"}],
        [{"sheet": "Data", "cell": "A1:B2", "value": 1}],
        [{"sheet": "Data", "cell": "XFE1", "value": 1}],
        [{"sheet": "Data", "cell": "A1", "formula": '=CALL("x")'}],
        [{"sheet": "", "cell": "A1", "value": 1}],
    ],
)
def test_validate_update_rejects(updates):
    assert _code(validate_update, UPDATES.validate_python(updates), []) == ErrorCode.VALIDATION_ERROR


def test_validate_update_requires_something():
    assert _code(validate_update, [], []) == ErrorCode.VALIDATION_ERROR


def test_update_unknown_sheet():
    updates = UPDATES.validate_python([{"sheet": "Missing", "cell": "A1", "value": 1}])
    assert _code(update_xlsx, FIXTURE.read_bytes(), updates, []) == ErrorCode.VALIDATION_ERROR


def test_update_warns_when_charts_are_dropped():
    wb = Workbook()
    ws = wb.active
    for v in (1, 2, 3):
        ws.append([v])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3))
    ws.add_chart(chart, "C1")
    buf = io.BytesIO()
    wb.save(buf)
    assert any(n.startswith("xl/charts/") for n in zipfile.ZipFile(buf).namelist())
    out, res = update_xlsx(buf.getvalue(), UPDATES.validate_python([{"sheet": "Sheet", "cell": "A1", "value": 9}]), [])
    assert any("chart" in w.lower() for w in res["warnings"])
